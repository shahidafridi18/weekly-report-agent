from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


NUMBER_FORMAT = '#,##0.00;[Red](#,##0.00);-'
PERCENT_FORMAT = '+0.00%;-0.00%;-'
HEADER_FILL = PatternFill("solid", fgColor="24364B")
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_BORDER = Border(
    bottom=Side(style="thin", color="20252A")
)
INSIGHT_HEADER_FILL = PatternFill("solid", fgColor="24364B")
INSIGHT_CONTENT_FILL = PatternFill("solid", fgColor="F3F7FA")


def _safe_cell_value(value):
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return f"'{value}"
    return value


def _write_table(
    workbook: Workbook,
    title: str,
    headers: list[str],
    rows: list[list],
    widths: list[int] | None = None,
    percentage_columns: set[int] | None = None,
) -> None:
    worksheet = workbook.create_sheet(title)
    worksheet.append(headers)

    for row in rows:
        worksheet.append([_safe_cell_value(value) for value in row])

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.page_margins.left = 0.25
    worksheet.page_margins.right = 0.25
    worksheet.oddFooter.center.text = "C2 - CONFIDENTIAL | Page &P of &N"

    # FCCR-inspired compact table styling
    header_fill = PatternFill("solid", fgColor="24364B")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    header_border = Border(
        bottom=Side(style="thin", color="24364B"),
        top=Side(style="thin", color="24364B"),
        left=Side(style="thin", color="24364B"),
        right=Side(style="thin", color="24364B"),
    )
    
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = header_border
        cell.alignment = header_alignment
    worksheet.row_dimensions[1].height = 32

    # Apply alternating row colors
    light_fill = PatternFill("solid", fgColor="F3F7FA")
    white_fill = PatternFill("solid", fgColor="FFFFFF")
    cell_border = Border(
        left=Side(style="thin", color="D4D8DC"),
        right=Side(style="thin", color="D4D8DC"),
        top=Side(style="thin", color="D4D8DC"),
        bottom=Side(style="thin", color="D4D8DC"),
    )
    cell_alignment = Alignment(vertical="center", wrap_text=False)

    for row_index, row in enumerate(worksheet.iter_rows(min_row=2), start=2):
        fill = light_fill if (row_index - 2) % 2 == 0 else white_fill
        for cell in row:
            cell.fill = fill
            cell.border = cell_border
            cell.alignment = cell_alignment

    # Set column widths
    for column_index in range(1, len(headers) + 1):
        if widths and column_index <= len(widths):
            width = widths[column_index - 1]
        else:
            width = min(
                max(
                    len(str(headers[column_index - 1])) + 2,
                    max(
                        (
                            len(str(worksheet.cell(row, column_index).value or ""))
                            for row in range(2, min(worksheet.max_row, 52) + 1)
                        ),
                        default=10,
                    ) + 2,
                ),
                42,
            )
        worksheet.column_dimensions[get_column_letter(column_index)].width = width

    # Format numbers and percentages
    percentage_columns = percentage_columns or set()
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = (
                    PERCENT_FORMAT
                    if cell.column in percentage_columns
                    else NUMBER_FORMAT
                )
                cell.alignment = Alignment(horizontal="right", vertical="center")


def _write_ai_insights(
    workbook: Workbook,
    ai_insights: str,
) -> None:
    """
    Create a dedicated sheet for AI insights
    with improved formatting and styling.
    """
    worksheet = workbook.create_sheet("AI Insights")
    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.oddHeader.left.text = "C2 - CONFIDENTIAL"
    worksheet.oddFooter.center.text = "C2 - CONFIDENTIAL | Page &P of &N"
    
    # Add title
    worksheet.merge_cells("A1:C1")
    title_cell = worksheet["A1"]
    title_cell.value = "AI-Generated Business Insights"
    title_cell.font = Font(bold=True, color="FFFFFF", size=14)
    title_cell.fill = INSIGHT_HEADER_FILL
    title_cell.alignment = Alignment(
        vertical="center",
        horizontal="left",
        wrap_text=True
    )
    worksheet.row_dimensions[1].height = 32
    
    worksheet.append([])  # Empty row for spacing
    
    # Parse insights with better formatting
    lines = ai_insights.split("\n")
    current_section = None
    section_fill = PatternFill("solid", fgColor="DCE8F2")
    
    for line in lines:
        line = line.strip()
        if not line:
            continue
            
        row_num = worksheet.max_row + 1
        
        # Handle markdown-style headers
        if line.startswith("## "):
            # Section header
            current_section = line[3:].strip()
            worksheet.merge_cells(f"A{row_num}:C{row_num}")
            cell = worksheet[f"A{row_num}"]
            cell.value = current_section
            cell.font = Font(bold=True, color="315F8C", size=12)
            cell.fill = PatternFill("solid", fgColor="DCE8F2")
            cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
            worksheet.row_dimensions[row_num].height = 24
            
        elif line.startswith("# "):
            # Main header
            main_header = line[2:].strip()
            worksheet.merge_cells(f"A{row_num}:C{row_num}")
            cell = worksheet[f"A{row_num}"]
            cell.value = main_header
            cell.font = Font(bold=True, color="24364B", size=13)
            cell.fill = PatternFill("solid", fgColor="DCE8F2")
            cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
            worksheet.row_dimensions[row_num].height = 26
            
        elif line.startswith("- "):
            # Bullet point
            bullet_text = line[2:].strip()
            worksheet.merge_cells(f"B{row_num}:C{row_num}")
            
            bullet_cell = worksheet[f"A{row_num}"]
            bullet_cell.value = "•"
            bullet_cell.font = Font(size=11, color="24364B", bold=True)
            bullet_cell.alignment = Alignment(horizontal="center")
            
            text_cell = worksheet[f"B{row_num}"]
            text_cell.value = bullet_text
            text_cell.font = Font(size=11, color="20252A")
            text_cell.fill = section_fill
            text_cell.alignment = Alignment(vertical="top", horizontal="left", wrap_text=True)
            
            lines_count = max(1, len(bullet_text) // 100 + 1)
            worksheet.row_dimensions[row_num].height = max(22, 16 * lines_count)
            
        elif line.startswith("**") or "**" in line:
            # Bold text
            text = line.replace("**", "")
            worksheet.merge_cells(f"A{row_num}:C{row_num}")
            cell = worksheet[f"A{row_num}"]
            cell.value = text
            cell.font = Font(bold=True, size=11, color="20252A")
            cell.fill = section_fill
            cell.alignment = Alignment(vertical="top", horizontal="left", wrap_text=True)
            lines_count = max(1, len(text) // 100 + 1)
            worksheet.row_dimensions[row_num].height = max(20, 16 * lines_count)
            
        else:
            # Regular text
            worksheet.merge_cells(f"A{row_num}:C{row_num}")
            cell = worksheet[f"A{row_num}"]
            cell.value = line
            cell.font = Font(size=11, color="495057")
            cell.fill = section_fill
            cell.alignment = Alignment(vertical="top", horizontal="left", wrap_text=True)
            lines_count = max(1, len(line) // 100 + 1)
            worksheet.row_dimensions[row_num].height = max(20, 16 * lines_count)
        
        # Add spacing row after each major element
        if line.startswith("## ") or line.startswith("# "):
            worksheet.append([])
    
    # Set column widths for better readability
    worksheet.column_dimensions["A"].width = 3
    worksheet.column_dimensions["B"].width = 60
    worksheet.column_dimensions["C"].width = 60


def generate_weekly_report_excel(
    analysis: dict,
    output_path: str | Path,
    previous_file_name: str,
    current_file_name: str,
    ai_insights: str | None = None,
) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    # Load existing workbook if it exists, otherwise create new one
    if output_path.exists():
        workbook = load_workbook(output_path)
        # Remove AI Insights sheet if it exists (will be recreated with new data)
        if "AI Insights" in workbook.sheetnames:
            workbook.remove(workbook["AI Insights"])
        # Skip recreating analysis sheets if they already exist
        # Just update AI Insights and save
        if ai_insights:
            _write_ai_insights(workbook, ai_insights)
        workbook.save(output_path)
        return output_path
    
    # Create new workbook if file doesn't exist
    workbook = Workbook()
    workbook.remove(workbook.active)

    columns = analysis.get("columns", {})
    key_columns = columns.get("key_columns", [])
    metric_columns = columns.get("numerical_columns", [])
    row_summary = analysis.get("row_summary", {})

    summary = workbook.create_sheet("Summary")
    summary.sheet_view.showGridLines = False
    summary.sheet_properties.pageSetUpPr.fitToPage = True
    summary.page_setup.fitToWidth = 1
    summary.page_setup.fitToHeight = 1
    summary.oddHeader.left.text = "C2 - CONFIDENTIAL"
    summary.oddHeader.right.text = "WEEKLY AI COMPARISON REPORT"
    summary.oddFooter.center.text = "C2 - CONFIDENTIAL | Page &P of &N"
    summary.merge_cells("A1:F1")
    summary["A1"] = "Weekly Business Comparison Report"
    summary["A1"].font = Font(bold=True, color="FFFFFF", size=16)
    summary["A1"].fill = PatternFill("solid", fgColor="24364B")
    summary["A1"].alignment = Alignment(vertical="center", horizontal="left")
    summary.row_dimensions[1].height = 30
    summary.append([])
    
    # Add summary information with better styling
    info_fill = PatternFill("solid", fgColor="DCE8F2")
    info_font = Font(bold=True, size=11, color="24364B")
    value_font = Font(size=11, color="495057")
    
    # Previous file row
    summary.append(["Previous file", previous_file_name])
    summary["A3"].font = info_font
    summary["A3"].fill = info_fill
    summary["B3"].font = value_font
    
    # Current file row
    summary.append(["Current file", current_file_name])
    summary["A4"].font = info_font
    summary["A4"].fill = info_fill
    summary["B4"].font = value_font
    
    # Matching keys row
    summary.append(["Matching keys", ", ".join(key_columns), "KPI fields", ", ".join(metric_columns)])
    summary["A5"].font = info_font
    summary["A5"].fill = info_fill
    summary["B5"].font = value_font
    summary["C5"].font = info_font
    summary["C5"].fill = info_fill
    summary["D5"].font = value_font
    
    summary.append([])
    
    # Summary statistics header row
    summary.append(["Previous rows", "Current rows", "Matched", "New", "Removed"])
    for col in ["A", "B", "C", "D", "E"]:
        summary[f"{col}7"].fill = PatternFill("solid", fgColor="24364B")
        summary[f"{col}7"].font = Font(bold=True, color="FFFFFF", size=11)
        summary[f"{col}7"].alignment = Alignment(horizontal="center", vertical="center")
    
    # Summary statistics data row
    summary.append(
        [
            row_summary.get("previous_rows", 0),
            row_summary.get("current_rows", 0),
            row_summary.get("matched_rows", 0),
            row_summary.get("new_rows", 0),
            row_summary.get("removed_rows", 0),
        ]
    )
    for col in ["A", "B", "C", "D", "E"]:
        summary[f"{col}8"].fill = PatternFill("solid", fgColor="F3F7FA")
        summary[f"{col}8"].font = Font(size=11, color="495057", bold=True)
        summary[f"{col}8"].alignment = Alignment(horizontal="center", vertical="center")
    
    # Set column widths
    for column, width in {"A": 20, "B": 30, "C": 16, "D": 16, "E": 16, "F": 16}.items():
        summary.column_dimensions[column].width = width

    # Add AI Insights sheet right after Summary for visibility
    if ai_insights:
        _write_ai_insights(
            workbook,
            ai_insights,
        )
        workbook._sheets.insert(1, workbook._sheets.pop())

    metric_summary = analysis.get("metric_summary", {})
    _write_table(
        workbook,
        "Metric Summary",
        ["Metric", "Previous", "Current", "Change", "Change %", "Direction"],
        [
            [
                metric,
                values.get("previous_total"),
                values.get("current_total"),
                values.get("absolute_change"),
                (values["percentage_change"] / 100)
                if values.get("percentage_change") is not None
                else None,
                values.get("direction", "").title(),
            ]
            for metric, values in metric_summary.items()
        ],
        widths=[38, 16, 16, 16, 14, 16],
        percentage_columns={5},
    )

    movement_bridge = analysis.get("movement_bridge", {})
    _write_table(
        workbook,
        "Movement Bridge",
        ["Metric", "Matched change", "New impact", "Removed impact", "Explained change", "Actual change"],
        [
            [
                metric,
                values.get("matched_entity_change"),
                values.get("new_entity_impact"),
                values.get("removed_entity_impact"),
                values.get("explained_change"),
                values.get("actual_total_change"),
            ]
            for metric, values in movement_bridge.items()
        ],
        widths=[38, 18, 16, 18, 18, 16],
    )

    variance_headers = [*key_columns, "Metric", "Previous", "Current", "Change", "Change %"]
    variance_rows = []
    for record in analysis.get("variance_data", []):
        entity = [record.get(key) for key in key_columns]
        for metric in metric_columns:
            percentage = record.get(f"{metric}_change_pct")
            variance_rows.append(
                entity
                + [
                    metric,
                    record.get(f"{metric}_previous"),
                    record.get(f"{metric}_current"),
                    record.get(f"{metric}_change"),
                    (percentage / 100) if percentage is not None else None,
                ]
            )
    _write_table(
        workbook,
        "Entity Variances",
        variance_headers,
        variance_rows,
        widths=[22] * len(key_columns) + [38, 16, 16, 16, 14],
        percentage_columns={len(variance_headers)},
    )

    new_headers = [*key_columns, "Metric", "Current value"]
    new_rows = [
        [item.get("entity", {}).get(key) for key in key_columns]
        + [metric, value]
        for item in analysis.get("new_entities", [])
        for metric, value in item.get("values", {}).items()
    ]
    _write_table(
        workbook,
        "New Entities",
        new_headers,
        new_rows,
        widths=[22] * len(key_columns) + [38, 18],
    )

    removed_headers = [*key_columns, "Metric", "Previous value"]
    removed_rows = [
        [item.get("entity", {}).get(key) for key in key_columns]
        + [metric, value]
        for item in analysis.get("removed_entities", [])
        for metric, value in item.get("previous_values", {}).items()
    ]
    _write_table(
        workbook,
        "Removed Entities",
        removed_headers,
        removed_rows,
        widths=[22] * len(key_columns) + [38, 18],
    )

    movement_headers = [*key_columns, "Metric", "Previous", "Current", "Change", "Change %", "Direction"]
    movement_rows = [
        [item.get("entity", {}).get(key) for key in key_columns]
        + [
            item.get("metric"),
            item.get("previous_value"),
            item.get("current_value"),
            item.get("absolute_change"),
            (item["percentage_change"] / 100)
            if item.get("percentage_change") is not None
            else None,
            item.get("direction"),
        ]
        for item in analysis.get("major_movements", [])
    ]
    _write_table(
        workbook,
        "Major Movements",
        movement_headers,
        movement_rows,
        widths=[22] * len(key_columns) + [38, 16, 16, 16, 14, 16],
        percentage_columns={len(movement_headers) - 1},
    )

    transition_headers = [*key_columns, "Metric", "Previous", "Current", "Transition"]
    transition_rows = [
        [item.get("entity", {}).get(key) for key in key_columns]
        + [
            item.get("metric"),
            item.get("previous_value"),
            item.get("current_value"),
            item.get("transition"),
        ]
        for item in analysis.get("zero_transitions", [])
    ]
    _write_table(
        workbook,
        "Zero Transitions",
        transition_headers,
        transition_rows,
        widths=[22] * len(key_columns) + [38, 16, 16, 20],
    )

    workbook.save(output_path)
    return output_path