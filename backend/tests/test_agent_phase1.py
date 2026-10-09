from pathlib import Path
import shutil
import json
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest
from unittest.mock import Mock, patch

from fastapi.testclient import TestClient
from google.genai import types
from openpyxl import load_workbook

from app.agent.models import CompareRequest, FileRequest, ReportRequest
from app.agent.orchestrator import AgentUnavailableError, GeminiToolSelector, WeeklyAgent
from app.agent.report_store import LocalReportStore
from app.agent.settings import AgentSettings
from app.agent.tools import WeeklyTools
from app.api.agent import get_agent
from app.ingestion.excel_reader import read_excel_bytes
from app.main import app
from app.processing.serializer import make_json_safe
from app.sources.base import AmbiguousFileError
from app.sources.local import LocalFileSource


FIXTURES = Path(__file__).resolve().parent / "data" / "weekly"


class Phase1Tests(unittest.TestCase):
    def setUp(self):
        self.temporary = TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.input = self.root / "weekly"
        self.input.mkdir()
        for week in (1, 2):
            filename = f"counterparty_week{week}_formatted.xlsx"
            shutil.copyfile(FIXTURES / filename, self.input / filename)
        self.settings = AgentSettings(self.input, self.root / "excel", self.root / "pdf")
        self.source = LocalFileSource(self.input)
        self.store = LocalReportStore(self.settings)
        self.tools = WeeklyTools(self.source, self.store)
        self.selector = Mock()
        self.agent = WeeklyAgent(self.tools, self.selector)
        app.dependency_overrides[get_agent] = lambda: self.agent
        self.addCleanup(app.dependency_overrides.clear)
        self.client = TestClient(app)
        self.addCleanup(self.client.close)

    def test_listing_excludes_lock_files_and_reports(self):
        for filename in ("~$week1.xlsx", "week1_vs_week2_report.xlsx", "notes.pdf"):
            (self.input / filename).write_bytes(b"not an input")
        result = self.client.get("/api/agent/files")
        self.assertEqual(result.status_code, 200)
        self.assertEqual([item["week_hint"] for item in result.json()["files"]], [1, 2])

    def test_week_selector_resolves_correct_file(self):
        info, content = self.source.read_file("week 2")
        self.assertEqual(info["filename"], "counterparty_week2_formatted.xlsx")
        self.assertEqual(read_excel_bytes(content).shape, (126, 28))

    def test_ambiguous_weeks_require_exact_selector(self):
        shutil.copyfile(self.input / "counterparty_week1_formatted.xlsx", self.input / "other_week1.xlsx")
        with self.assertRaises(AmbiguousFileError):
            self.source.read_file("week 1")
        result = self.client.post("/api/agent/analyze", json={"file": "week 1"})
        self.assertEqual(result.status_code, 409)
        self.assertEqual(len(result.json()["detail"]["candidates"]), 2)

    def test_paths_cannot_escape_input_root(self):
        for selector in ("../outside.xlsx", str(self.root / "outside.xlsx")):
            with self.assertRaises(ValueError):
                self.source.read_file(selector)

    def test_symlink_outside_input_root_is_not_listed(self):
        outside = self.root / "outside.xlsx"
        outside.write_bytes(b"not a workbook")
        try:
            (self.input / "outside_week9.xlsx").symlink_to(outside)
        except OSError:
            self.skipTest("OS does not permit symlink creation")
        self.assertEqual(len(self.source.list_files()), 2)

    def test_input_size_limit(self):
        with self.assertRaisesRegex(ValueError, "size limit"):
            LocalFileSource(self.input, max_bytes=1).read_file("week 1")

    def test_input_report_folders_cannot_overlap(self):
        with self.assertRaises(ValueError):
            AgentSettings(self.input, self.input / "reports", self.root / "pdf")

    def test_single_week_returns_validated_numeric_statistics(self):
        result = self.client.post("/api/agent/analyze", json={"file": "week 1"})
        self.assertEqual(result.status_code, 200)
        data = result.json()
        self.assertEqual((data["row_count"], data["column_count"]), (126, 28))
        self.assertEqual(len(data["metric_statistics"]), 24)
        self.assertNotIn("Rank", data["metric_statistics"])
        self.assertIs(data["validation"]["valid"], True)

    def test_existing_pipeline_runs_from_local_folder(self):
        result = self.client.post("/api/agent/compare", json={"previous_file": "week 1", "current_file": "week 2", "metrics": ["Gross CE"]})
        self.assertEqual(result.status_code, 200)
        data = result.json()
        self.assertEqual(data["analysis"]["row_summary"], {
            "previous_rows": 126, "current_rows": 126, "matched_rows": 125, "new_rows": 1, "removed_rows": 1,
        })
        self.assertEqual(list(data["selected_metric_summary"]), ["Gross CE"])
        self.assertEqual(len(data["analysis"]["metric_summary"]), 24)

    def test_same_file_and_unknown_metric_rejected(self):
        for body in (
            {"previous_file": "week 1", "current_file": "week 1"},
            {"previous_file": "week 1", "current_file": "week 2", "metrics": ["Invented Metric"]},
        ):
            self.assertEqual(self.client.post("/api/agent/compare", json=body).status_code, 400)

    def test_duplicate_business_keys_still_fail_validation(self):
        _, content = self.source.read_file("week 1")
        frame = read_excel_bytes(content)
        frame.loc[1, ["SIREN", "Unique Identifier"]] = frame.loc[0, ["SIREN", "Unique Identifier"]].values
        with patch("app.agent.tools.read_excel_bytes", return_value=frame):
            result = self.client.post("/api/agent/compare", json={"previous_file": "week 1", "current_file": "week 2"})
        self.assertEqual(result.status_code, 400)
        self.assertIn("duplicate business keys", result.json()["detail"])

    def test_ai_failure_does_not_break_comparison(self):
        with patch("app.insights.llm_service.GeminiService", side_effect=RuntimeError("Unavailable")):
            result = self.tools.compare_weekly_files(CompareRequest(previous_file="week 1", current_file="week 2", generate_ai_insights=True))
        self.assertEqual(result["ai"]["status"], "unavailable")
        self.assertEqual(result["analysis"]["row_summary"]["matched_rows"], 125)

    def test_reports_persist_in_separate_folders_and_can_be_downloaded(self):
        result = self.client.post("/api/agent/reports", json={"previous_file": "week 1", "current_file": "week 2", "report_format": "both"})
        self.assertEqual(result.status_code, 200)
        reports = result.json()["reports"]
        self.assertEqual({item["format"] for item in reports}, {"pdf", "xlsx"})
        for report in reports:
            path = self.store.get_path(report["format"], report["report_id"])
            self.assertEqual(path.parent, self.root / ("excel" if report["format"] == "xlsx" else "pdf"))
            downloaded = self.client.get(report["download_url"])
            self.assertEqual(downloaded.status_code, 200)
            self.assertEqual(downloaded.content, path.read_bytes())
            self.assertTrue(path.exists(), "Downloads must not delete persistent outputs")
            if report["format"] == "xlsx":
                workbook = load_workbook(path, read_only=True, data_only=True)
                self.assertEqual(workbook["Entity Variances"].max_row, 3001)
                workbook.close()
            else:
                self.assertTrue(path.read_bytes().startswith(b"%PDF"))
        self.assertEqual(len(self.source.list_files()), 2)

    def test_report_runs_use_unique_output_ids(self):
        request = ReportRequest(previous_file="week 1", current_file="week 2", report_format="xlsx")
        first = self.tools.generate_weekly_reports(request)["reports"][0]
        second = self.tools.generate_weekly_reports(request)["reports"][0]
        self.assertNotEqual(first["report_id"], second["report_id"])

    def test_chat_uses_selected_tool_and_python_facts(self):
        self.selector.select.return_value = ("compare_weekly_files", {"previous_file": "week 1", "current_file": "week 2", "metrics": ["Gross CE"]})
        result = self.client.post("/api/agent/chat", json={"message": "Compare week 1 and week 2 Gross CE"})
        self.assertEqual(result.status_code, 200)
        self.assertEqual(result.json()["status"], "success")
        self.assertIn("125 matched", result.json()["answer"])
        self.assertIn("Gross CE", result.json()["answer"])

    def test_chat_missing_or_ambiguous_file_returns_choices(self):
        shutil.copyfile(self.input / "counterparty_week1_formatted.xlsx", self.input / "other_week1.xlsx")
        for selector in ("week 1", "week 3"):
            self.selector.select.return_value = ("analyze_weekly_file", {"file": selector})
            result = self.agent.chat("Analyze " + selector)
            self.assertEqual(result["status"], "needs_clarification")

    def test_chat_rejects_unknown_tool_and_invalid_parameters(self):
        self.selector.select.return_value = ("run_shell", {"command": "ls"})
        self.assertEqual(self.client.post("/api/agent/chat", json={"message": "List inputs"}).status_code, 400)
        self.selector.select.return_value = ("analyze_weekly_file", {"file": "week 1", "unknown": True})
        self.assertEqual(self.client.post("/api/agent/chat", json={"message": "Analyze week 1"}).status_code, 503)

    def test_chat_sdk_transport_and_tool_schema(self):
        llm = Mock()
        llm.model = "configured-working-model"
        llm.client.models.generate_content.return_value = types.GenerateContentResponse(
            candidates=[types.Candidate(content=types.Content(role="model", parts=[types.Part(function_call=types.FunctionCall(name="analyze_weekly_file", args={"file": "week 1"}))]))]
        )
        with patch("app.insights.llm_service.GeminiService", return_value=llm):
            selected = GeminiToolSelector().select("Analyze week 1", self.source.list_files())
        self.assertEqual(selected, ("analyze_weekly_file", {"file": "week 1"}))
        kwargs = llm.client.models.generate_content.call_args.kwargs
        self.assertEqual(kwargs["model"], llm.model)
        self.assertEqual(kwargs["config"].tool_config.function_calling_config.mode, types.FunctionCallingConfigMode.ANY)
        self.assertNotIn("Counterparty 1", kwargs["contents"])

    def test_chat_provider_failure_and_multiple_calls_are_not_executed(self):
        llm = Mock()
        with patch("app.insights.llm_service.GeminiService", side_effect=RuntimeError("missing key")):
            with self.assertRaises(AgentUnavailableError):
                GeminiToolSelector().select("List inputs", self.source.list_files())
        llm.client.models.generate_content.return_value = SimpleNamespace(function_calls=[Mock(), Mock()])
        with patch("app.insights.llm_service.GeminiService", return_value=llm):
            with self.assertRaises(AgentUnavailableError):
                GeminiToolSelector().select("List inputs", self.source.list_files())

    def test_validation_rejects_blank_message_and_negative_threshold(self):
        self.assertEqual(self.client.post("/api/agent/chat", json={"message": "   "}).status_code, 422)
        self.assertEqual(self.client.post("/api/agent/compare", json={"previous_file": "week 1", "current_file": "week 2", "movement_threshold_pct": -1}).status_code, 422)

    def test_original_comparison_endpoint_still_works(self):
        with (self.input / "counterparty_week1_formatted.xlsx").open("rb") as previous, (self.input / "counterparty_week2_formatted.xlsx").open("rb") as current:
            result = self.client.post("/api/compare", files={"previous_file": ("week1.xlsx", previous), "current_file": ("week2.xlsx", current)}, data={"generate_ai_insights": "false"})
        self.assertEqual(result.status_code, 200)
        self.assertEqual(result.json()["analysis"]["row_summary"]["matched_rows"], 125)

    def test_nested_nonfinite_values_become_json_null(self):
        result = make_json_safe({"valid": True, "nested": [{"missing": float("nan"), "infinite": float("inf")} ]})
        self.assertIs(result["valid"], True)
        self.assertEqual(result["nested"], [{"missing": None, "infinite": None}])
        json.dumps(result, allow_nan=False)

    def test_original_report_endpoint_still_works(self):
        with (self.input / "counterparty_week1_formatted.xlsx").open("rb") as previous, (self.input / "counterparty_week2_formatted.xlsx").open("rb") as current:
            result = self.client.post("/api/report", files={"previous_file": ("week1.xlsx", previous), "current_file": ("week2.xlsx", current)}, data={"generate_ai_insights": "false", "report_format": "pdf"})
        self.assertEqual(result.status_code, 200)
        self.assertTrue(result.content.startswith(b"%PDF"))


if __name__ == "__main__":
    unittest.main()
