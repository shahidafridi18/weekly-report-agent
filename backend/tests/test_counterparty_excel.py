from pathlib import Path
import re
from tempfile import TemporaryDirectory
import unittest

from openpyxl import load_workbook

from app.analysis.engine import run_analysis
from app.ingestion.excel_reader import (
    read_excel_bytes,
    read_excel_file,
)
from app.schema import COUNTERPARTY_METRIC_COLUMNS
from app.reporting.chart_generator import (
    generate_kpi_change_chart,
    generate_movement_bridge_chart,
)
from app.reporting.excel_report_generator import (
    generate_weekly_report_excel,
)
from app.reporting.pdf_generator import generate_weekly_report_pdf


SAMPLE_DATA_DIR = Path(__file__).resolve().parents[1] / "sample_data"
KEY_COLUMNS = ["SIREN", "Unique Identifier"]


class CounterpartyExcelTests(unittest.TestCase):
    def read_week(self, week: int):
        path = SAMPLE_DATA_DIR / f"counterparty_week{week}_formatted.xlsx"
        return read_excel_file(path)

    def test_reads_flat_analysis_sheet_and_preserves_identifiers(self):
        frame = self.read_week(1)

        self.assertEqual(frame.shape, (126, 28))
        self.assertEqual(str(frame["SIREN"].dtype), "string")
        self.assertEqual(str(frame["Unique Identifier"].dtype), "string")
        self.assertNotIn("Rank", COUNTERPARTY_METRIC_COLUMNS)

    def test_reads_grouped_template_and_removes_total_row(self):
        path = SAMPLE_DATA_DIR / "counterparty_week1_formatted.xlsx"
        frame = read_excel_bytes(
            path.read_bytes(),
            sheet_name="FSBII - Etat A - Hebdo Regulate",
        )

        self.assertEqual(frame.shape, (126, 28))
        self.assertEqual(frame.iloc[-1]["Counterparty Name"], "Counterparty 126")
        self.assertNotIn("TOTAL EXPOSURE", frame["Counterparty Name"].tolist())

    def test_weekly_analysis_uses_exposure_metrics_and_detects_churn(self):
        week1 = self.read_week(1)
        week2 = self.read_week(2)
        result = run_analysis(week1, week2, KEY_COLUMNS)

        self.assertEqual(
            result["columns"]["numerical_columns"],
            list(COUNTERPARTY_METRIC_COLUMNS),
        )
        self.assertEqual(
            result["row_summary"],
            {
                "previous_rows": 126,
                "current_rows": 126,
                "matched_rows": 125,
                "new_rows": 1,
                "removed_rows": 1,
            },
        )

    def test_rejects_duplicate_business_keys(self):
        week1 = self.read_week(1)
        week2 = self.read_week(2)
        week1.loc[1, KEY_COLUMNS] = week1.loc[0, KEY_COLUMNS].values

        with self.assertRaisesRegex(ValueError, "duplicate business keys"):
            run_analysis(week1, week2, KEY_COLUMNS)

    def test_rejects_missing_counterparty_metric(self):
        week1 = self.read_week(1).drop(columns=["Gross CE"])
        week2 = self.read_week(2).drop(columns=["Gross CE"])

        with self.assertRaisesRegex(ValueError, "required counterparty metrics"):
            run_analysis(week1, week2, KEY_COLUMNS)

    def test_legacy_product_workbook_still_reads(self):
        frame = read_excel_file(SAMPLE_DATA_DIR / "week1.xlsx")

        self.assertEqual(len(frame), 4)
        self.assertIn("Product", frame.columns)
        self.assertIn("Revenue", frame.columns)

    def test_excel_report_contains_filterable_full_variances(self):
        week1 = self.read_week(1)
        week2 = self.read_week(2)
        analysis = run_analysis(week1, week2, KEY_COLUMNS)

        with TemporaryDirectory() as directory:
            report_path = generate_weekly_report_excel(
                analysis,
                Path(directory) / "comparison.xlsx",
                "week1.xlsx",
                "week2.xlsx",
            )
            workbook = load_workbook(
                report_path,
                data_only=True,
            )
            try:
                self.assertIn("Metric Summary", workbook.sheetnames)
                self.assertIn("Entity Variances", workbook.sheetnames)
                self.assertEqual(
                    workbook["Entity Variances"].max_row,
                    1 + 125 * len(COUNTERPARTY_METRIC_COLUMNS),
                )
                self.assertEqual(
                    workbook["Entity Variances"].auto_filter.ref,
                    "A1:G3001",
                )
            finally:
                workbook.close()

    def test_pdf_summary_omits_full_variance_appendix(self):
        week1 = self.read_week(1)
        week2 = self.read_week(2)
        analysis = run_analysis(week1, week2, KEY_COLUMNS)

        with TemporaryDirectory() as directory:
            directory = Path(directory)
            chart_path = generate_kpi_change_chart(
                analysis["metric_summary"],
                directory / "kpi.png",
            )
            bridge_path = generate_movement_bridge_chart(
                analysis["movement_bridge"],
                directory / "bridge.png",
            )
            report_path = generate_weekly_report_pdf(
                analysis=analysis,
                ai_insights=None,
                chart_path=chart_path,
                movement_bridge_chart_path=bridge_path,
                output_path=directory / "comparison.pdf",
                previous_file_name="week1.xlsx",
                current_file_name="week2.xlsx",
            )

            pdf_bytes = report_path.read_bytes()
            page_count = len(re.findall(rb"/Type\s*/Page\b", pdf_bytes))
            self.assertTrue(pdf_bytes.startswith(b"%PDF"))
            self.assertLess(page_count, 20)
            self.assertNotIn(b"Detailed Variance Appendix", pdf_bytes)


if __name__ == "__main__":
    unittest.main()